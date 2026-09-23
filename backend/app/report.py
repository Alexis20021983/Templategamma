from io import BytesIO
from docx import Document
from docx.shared import Inches
from pathlib import Path

def build_report(case) -> BytesIO:
    doc = Document()
    doc.add_heading(f"QA Evidence Report: {case.title}", 0)
    doc.add_paragraph(case.description or "No description")
    doc.add_paragraph(f"Status: {case.status}   Owner: {case.owner or 'Unassigned'}")
    for label in ("requirement", "objective", "preconditions", "test_data", "tester", "reviewer"):
        value = getattr(case, label, "")
        if value: doc.add_paragraph(f"{label.replace('_', ' ').title()}: {value}")
    for index, step in enumerate(case.steps, 1):
        doc.add_heading(f"{index}. {step.title}", level=1)
        doc.add_paragraph(f"Status: {step.status}")
        doc.add_heading("Expected result", level=2)
        doc.add_paragraph(step.expected_result or "—")
        doc.add_heading("Actual result", level=2)
        doc.add_paragraph(step.actual_result or "—")
        if step.evidences:
            doc.add_heading("Evidence", level=2)
            for evidence in step.evidences:
                doc.add_paragraph(f"{evidence.name}: {evidence.url or '(no URL)'}", style="List Bullet")
                if evidence.notes:
                    doc.add_paragraph(evidence.notes)
                if evidence.file_path and Path(evidence.file_path).exists():
                    try:
                        doc.add_picture(evidence.file_path, width=Inches(5.8))
                    except Exception:
                        doc.add_paragraph("[Image could not be embedded]")
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
